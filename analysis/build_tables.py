"""Tezdeki çizelgeleri veriden yeniden üretir ve belgedekiyle karşılaştırır.

Çizelgeler bugüne kadar elle bakımda tutuldu: bir koşum yenilendiğinde ya da bir
kol eklendiğinde sayıların Word belgesine tek tek taşınması gerekti. Bu, sessiz
bir kayma kaynağıdır; belgedeki bir sayı veriden koptuğunda hiçbir şey uyarmaz.

Bu betik iki iş yapar:

  üret        Veriden türetilebilen her çizelgeyi yeniden kurar ve Word'e
              yapıştırılmaya hazır sekme ayraçlı dosyalar yazar.
  karşılaştır Ürettiğini belgedeki çizelgeyle hücre hücre karşılaştırır ve
              farkları listeler. Asıl değeri budur: belgenin veriden kaydığı
              yeri gösterir.

**Sonuç dosyalarına yazmaz.** Yalnız `--out` ile verilen klasöre yazar; varsayılan
`ieee/cizelgeler/`. `federated_evaluation_reports/`, `model_pickle/`,
`generated_visuals/` ve benzerleri salt okunur kullanılır.

Elle bakımda kalan çizelgeler ve sebepleri `ELLE` sözlüğünde yazılıdır; bunlar
veriden değil karardan doğar ve betik onlara dokunmaz.

    python analysis/build_tables.py
    python analysis/build_tables.py --only 5.3,5.8
    python analysis/build_tables.py --check          yalnız karşılaştır, yazma
    python analysis/build_tables.py --out /tmp/ciz
"""

import argparse
import glob
import os
import pickle
import re
import sys

import numpy as np
import pandas as pd

# Çizelgeler U+2212 eksi işareti taşır ve Windows konsolu varsayılan olarak
# cp1254 konuşur; yeniden yapılandırılmazsa betik yazdırırken çöker.
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, KOK)
sys.path.insert(0, os.path.join(KOK, "analysis"))

OZET = "experiment_comparison_summary.csv"
BELGE = "ieee/23435004022_duzeltilmis.docx"
VARSAYILAN_CIKTI = "ieee/cizelgeler"

# Çizelgeler ne kadar sayı taşırsa taşısın, bir kısmı ölçümden değil karardan
# doğar. Bunları üretmek, olmayan bir kaynağı varmış gibi göstermek olurdu.
ELLE = {
    "2.1": "literatür konumlandırması; üç niteliğe göre okuma kararı",
    "4.2": "açıklama sütunu elle yazılmış (ölçülen sütunları 4.2m üretir)",
    "5.1": "konfigürasyon listesi; her satır bir tasarım sorusu",
    "5.2": "prototip dönemi tarihçesi; yeniden koşulamaz",
    "5.11": "üst bloğu yayımlanmış literatür değeri, veriden türemez",
}


def tr(x, basamak=4):
    """Ondalık ayracı virgül. Tezin tamamı bu biçimde."""
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "-"
    return f"{x:.{basamak}f}".replace(".", ",")


def tri(x, basamak=4):
    """İşaretli, Türkçe ondalık. Eksi işareti U+2212; tez baştan sona böyle."""
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "-"
    govde = f"{abs(x):.{basamak}f}".replace(".", ",")
    return ("+" if x >= 0 else "−") + govde


def yuzde(x, basamak=1):
    return f"%{x:.{basamak}f}".replace(".", ",")


def _ozet():
    """Özet dosyası, plato penceresi yapılandırmayla uyuşuyorsa.

    Uyuşmadığında dosya bayattır ve PR-AUC sütunu yanlış bir pencereyle
    hesaplanmıştır. Bu sessizce geçilirse çizelgeler veriden değil eski bir
    koşumdan üretilir; betiğin ilk sürümü tam olarak buna düştü ve baseline'ı
    0,8371 yerine 0,8857 gösterdi. Yalnız pencereye bağlı olmayan sütunlar
    (iletişim baytları, yakınsama turu) bu durumda da kullanılabilir.
    """
    yol = os.path.join(KOK, OZET)
    if not os.path.exists(yol):
        raise SystemExit(f"{OZET} bulunamadı. Önce compare_experiments.py koşulmalı.")
    d = pd.read_csv(yol)
    from config_manager import config
    istenen = config.get("evaluation", "plateau_window_rounds")
    goruleni = sorted(d["Plateau_Rounds"].dropna().unique())
    if goruleni and float(istenen) not in [float(x) for x in goruleni]:
        d.attrs["bayat"] = (f"özet dosyası {goruleni} turluk plato penceresiyle "
                            f"üretilmiş, yapılandırma {istenen} istiyor")
    return d


_PRAUC = {}
_KOSUMLAR = {}


def _kosumlar():
    """Diskteki koşumlar, bir kez okunur.

    `load_runs` her çağrıda sürüm filtresinin atladığı kolları yazdırır; kol
    başına çağrılırsa aynı uyarı onlarca kez tekrarlanır ve gerçek uyarıları
    bastırır.
    """
    if not _KOSUMLAR:
        from seed_variance import load_runs
        _KOSUMLAR.update(load_runs())
    return _KOSUMLAR


def _plato_prauc(ad):
    """Bir kolun seed başına plato PR-AUC değerleri, canlı yapılandırmadan.

    Özet dosyasından değil diskteki tur sonuçlarından hesaplanır, çünkü tezin
    raporladığı istatistik yapılandırmadaki plato penceresine bağlıdır ve özet
    dosyası o pencere değiştiğinde bayatlar.
    """
    if ad in _PRAUC:
        return _PRAUC[ad]
    kosumlar = _kosumlar()
    if ad not in kosumlar:
        raise SystemExit(f"'{ad}' için tamamlanmış koşum yok.")
    from seed_variance import plateau_frames, plateau_pr_auc
    degerler = []
    for _, dizin in sorted(kosumlar[ad], key=lambda x: x[1]):
        cerceve = plateau_frames(dizin)
        if cerceve:
            degerler.append(plateau_pr_auc(cerceve))
    if not degerler:
        raise SystemExit(f"'{ad}' puanlanmış tur taşımıyor.")
    _PRAUC[ad] = np.asarray(degerler, dtype=float)
    return _PRAUC[ad]


def _kol(df, ad):
    """Bir deneyin seed'ler üzerindeki satırları. Boşsa açıkça söyler."""
    alt = df[df["Experiment"] == ad]
    if alt.empty:
        raise SystemExit(f"'{ad}' özet dosyasında yok. Koşum eksik olabilir.")
    return alt


def _ort(df, ad, sutun):
    return float(_kol(df, ad)[sutun].mean())


def _sd(df, ad, sutun):
    return float(_kol(df, ad)[sutun].std(ddof=1))


def _pr(ad):
    return float(_plato_prauc(ad).mean())


def _prsd(ad):
    return float(_plato_prauc(ad).std(ddof=1))


# --- Çizelge 4.1: öznitelik kümesinin kanallara dağılımı --------------------

KANALLAR = [
    ("Zaman bağlamı ve genel etkinlik",
     lambda f: f in ("isweekday", "isweekend") or "allact" in f),
    ("Oturum (logon)", lambda f: f.startswith("logon") or "logon" in f),
    ("Çıkarılabilir cihaz (USB)", lambda f: "usb" in f),
    ("Dosya (file)", lambda f: f.startswith("file") or "file" in f),
    ("E-posta (email)", lambda f: f.startswith("email") or "email" in f),
    ("Web (http)", lambda f: f.startswith("http") or "http" in f),
]


def cizelge_4_1():
    """Kanal başına öznitelik sayısı.

    Belgedeki üçüncü sütun (örnekler) elle seçilmiştir: kanalı en iyi anlatan
    üç ad, alfabetik ya da konumsal ilk üç değil. Üretilmez, karşılaştırmada da
    dikkate alınmaz.
    """
    from config_manager import config
    ozn = list(config.get("data", "selected_features"))
    kalan = list(ozn)
    satirlar = []
    for ad, kural in KANALLAR:
        bu = [f for f in kalan if kural(f)]
        kalan = [f for f in kalan if f not in bu]
        satirlar.append([ad, str(len(bu))])
    if kalan:
        raise SystemExit(f"Kanala düşmeyen öznitelik: {kalan}")
    satirlar.append(["Toplam", str(len(ozn))])
    return ["Aktivite kanalı", "Öznitelik"], satirlar


# --- Çizelge 4.2m: ölçülen sütunlar (std ve sıfır oranı) --------------------

def cizelge_4_2m(veri=None):
    from config_manager import config
    ozn = list(config.get("data", "selected_features"))
    yol = veri or os.path.join(KOK, "ExtractedData", "dayr4.2-percentile30.csv")
    if not os.path.exists(yol):
        raise SystemExit(f"{yol} yok; 4.2'nin ölçülen sütunları üretilemez.")
    d = pd.read_csv(yol, usecols=ozn)
    satirlar = []
    for f in ozn:
        s = d[f]
        # ddof=1. Belgedeki sütun `analyze_features.py`'den geldi ve pandas'ın
        # varsayılanını kullanıyor; `feature_selection_rule.py` ile
        # `task._drop_constant_features` de öyle. Burada ddof=0 kullanmak
        # 308.023 satırda dördüncü ondalığı on bir satırda kaydırıyor ve
        # `--check` bunları belgenin hatası gibi gösteriyordu. Referans hata
        # sapması σⱼ (`task.py`, Denklem 4.8) ayrı bir nicelik ve orada ddof=0
        # bilinçlidir; bu iki sapma birbirine çevrilmez.
        satirlar.append([f, tr(float(s.std(ddof=1)), 4),
                         tr(100.0 * float((s == 0).mean()), 1)])
    return ["Öznitelik", "Std", "Sıfır %"], satirlar


def _yukleme_mesaji():
    """Bir koşumun ürettiği yükleme mesajı sayısı: tur × örneklenen istemci.

    İstemci sayısı `[tool.fueba]` altında değil, Flower'ın kendi federasyon
    tanımındadır (`tool.flwr...options.num-supernodes`), o yüzden doğrudan
    `pyproject.toml`'dan okunur.
    """
    from config_manager import config
    try:
        import tomllib
    except ImportError:
        import tomli as tomllib
    with open(os.path.join(KOK, "pyproject.toml"), "rb") as f:
        ham = tomllib.load(f)
    fed = ham.get("tool", {}).get("flwr", {}).get("federations", {})
    istemci = None
    for ad, govde in fed.items():
        if isinstance(govde, dict) and "options" in govde:
            istemci = govde["options"].get("num-supernodes", istemci)
    if istemci is None:
        raise SystemExit("num-supernodes pyproject.toml'da bulunamadı.")
    tur = config.get("federation", "num_rounds")
    oran = config.get("federation", "fraction_fit")
    return tur * round(istemci * oran)


# --- Çizelge 4.3: analitik boyut ile telde ölçülen boyut --------------------

def cizelge_4_3():
    from config_manager import config
    from federated_ueba.task import Architecture, LSTMAutoencoder
    arch = Architecture.from_config(config)
    n = len(config.get("data", "selected_features"))
    model = LSTMAutoencoder(n, arch=arch)
    par = sum(p.numel() for p in model.parameters() if p.requires_grad)
    analitik = par * 4 / 1024 / 1024

    # Tek bir checkpoint değil, koşum boyunca yüklenen mesajların ortalaması.
    # Bir checkpoint'in kodlanmış boyutu tura göre değişir (ağırlıklar
    # oturdukça entropi düşer); tezin raporladığı büyüklük ortalamadır.
    d = _ozet()
    mesaj = _yukleme_mesaji()
    olculen = _ort(d, "baseline", "Upload_MB") / mesaj
    return (["Ölçüm", "Değer"],
            [[f"Analitik yoğun fp32 boyutu ({par:,}".replace(",", ".") + " × 4 B)",
              f"{tr(analitik, 4)} MB"],
             ["Ölçülen ortalama yükleme mesajı", f"{tr(olculen, 4)} MB"],
             ["Kayıpsız kodlamanın katkısı", yuzde(100 * (1 - olculen / analitik))]])


def _checkpoint(kol, seed, tur=None):
    kalip = os.path.join(KOK, "model_pickle", f"{kol}__seed{seed}",
                         "parameters_round_*.pkl")
    yollar = glob.glob(kalip)
    if not yollar:
        raise SystemExit(f"{kol}__seed{seed} için checkpoint yok.")
    def no(p):
        m = re.search(r"parameters_round_(\d+)", p)
        return int(m.group(1)) if m else -1
    yol = (max(yollar, key=no) if tur is None
           else next(p for p in yollar if no(p) == tur))
    with open(yol, "rb") as f:
        icerik = pickle.load(f)
    return icerik["global_parameters"] if isinstance(icerik, dict) else icerik


# --- Çizelge 5.3: iletişim ekseni ------------------------------------------

ILETISIM_SIRA = ["baseline", "quantization-fp16", "bidirectional-fp16",
                 "delta-0.1", "delta-0.05", "delta-0.05-downlink-fp16",
                 "top-k-0.1", "top-k-0.1-quant-fp16", "top-k-0.05"]


def cizelge_5_3():
    d = _ozet()
    temel = _ort(d, "baseline", "Total_Comm_MB")
    satirlar = []
    for ad in ILETISIM_SIRA:
        toplam = _ort(d, ad, "Total_Comm_MB")
        satirlar.append([
            ad, tr(_pr(ad)), tr(_prsd(ad)),
            tr(_ort(d, ad, "Upload_MB"), 1), tr(_ort(d, ad, "Download_MB"), 1),
            tr(toplam, 1),
            "referans" if ad == "baseline" else yuzde(100 * (1 - toplam / temel)),
        ])
    return (["Deney", "PR-AUC", "sd", "Yükleme (MB)", "İndirme (MB)",
             "Toplam (MB)", "Tasarruf"], satirlar)


# --- Çizelge 5.4: yayınlanan küresel modelin seyrekliği ---------------------

SEYREKLIK = [("baseline", "baseline"),
             ("top-k-0.1", "Ağırlık seyrekleştirme, α = 0,10"),
             ("top-k-0.05", "Ağırlık seyrekleştirme, α = 0,05"),
             ("delta-0.1", "Güncelleme seyrekleştirme, α = 0,10"),
             ("delta-0.05", "Güncelleme seyrekleştirme, α = 0,05")]


def cizelge_5_4(seed=1):
    satirlar = []
    for kol, etiket in SEYREKLIK:
        w = _checkpoint(kol, seed)
        duz = np.concatenate([a.ravel() for a in w])
        satirlar.append([etiket, tr(float(np.linalg.norm(duz)), 2),
                         yuzde(100.0 * float((duz == 0).mean()))])
    return ["Koşum", "Model normu", "Tam sıfır ağırlık"], satirlar


# --- Çizelge 5.5: yakınsama ve plato turu ----------------------------------

def cizelge_5_5():
    d = _ozet()
    satirlar = []
    for ad in ILETISIM_SIRA:
        k = _kol(d, ad)
        y, p = k["Conv_Round"], k["Plateau_Round"]
        satirlar.append([ad, tr(float(y.mean()), 1),
                         f"{int(y.min())}-{int(y.max())}",
                         tr(float(p.mean()), 1),
                         f"{int(p.min())}-{int(p.max())}"])
    return ["Deney", "Yakınsama turu", "Aralık", "Plato turu", "Aralık"], satirlar


# --- Çizelge 5.6 ve 5.7: ablasyonlar ---------------------------------------

def _fark(a, b, yineleme, guven):
    from seed_variance import plateau_frames, paired_difference
    kosumlar = _kosumlar()
    def kol(ad):
        return [f for f in (plateau_frames(dz)
                            for _, dz in sorted(kosumlar.get(ad, []),
                                                key=lambda x: x[1])) if f]
    ka, kb = kol(a), kol(b)
    if not ka or not kb:
        return None
    return paired_difference(ka, kb, yineleme, guven)


def _bootstrap_ayari():
    """Yineleme ve güven, `findings.py` ile aynı kaynaktan.

    Bootstrap aralıkları yineleme sayısına duyarlıdır; farklı bir sayıyla
    koşmak tezdekiyle uyuşmayan uç noktalar üretir ve bu, gerçek bir kayma
    sanılabilir.
    """
    from config_manager import config
    return (config.get("evaluation", "bootstrap_iterations"),
            config.get("evaluation", "bootstrap_confidence"))


def cizelge_5_6(yineleme=None, guven=None):
    if yineleme is None or guven is None:
        yineleme, guven = _bootstrap_ayari()
    d = _ozet()
    satirlar = [["baseline", tr(_pr("baseline")),
                 tr(_prsd("baseline")), "referans", "-"]]
    for kol, karar in [("no-bottleneck", "Bileşen gerekli"),
                       ("encoder-unidirectional", "Ayırt edilemiyor"),
                       ("encoder-unidirectional-64", "Ayırt edilemiyor"),
                       ("features-filtered", "Aralık geniş")]:
        f = _fark("baseline", kol, yineleme, guven)
        aralik = ("-" if f is None
                  else f"[{tri(f['CI_Lo'])}, {tri(f['CI_Hi'])}]")
        satirlar.append([kol, tr(_pr(kol)),
                         tr(_prsd(kol)), aralik, karar])
    return ["Konfigürasyon", "PR-AUC", "sd", "Katkının %95 aralığı", "Karar"], satirlar


def cizelge_5_7(yineleme=None, guven=None):
    if yineleme is None or guven is None:
        yineleme, guven = _bootstrap_ayari()
    from findings import ABLATIONS
    d = _ozet()
    satirlar = [["Üç aşamalı hat (nihai)", tr(_pr("ablation-full")),
                 "referans", "-", "-"]]
    etiket = {"persistence": "Zamansal kalıcılık yok",
              "zscore": "Z-skor kalibrasyonu yok",
              "topk": "Top-m öznitelik odağı yok",
              "diversity": "Çeşitlilik çarpanı eklenmiş"}
    sira = ["persistence", "zscore", "topk", "diversity"]
    kayit = {a[0]: a for a in ABLATIONS}
    for anahtar in sira:
        _, _, olan, olmayan = kayit[anahtar]
        yok = olmayan if anahtar != "diversity" else olan
        f = _fark(olan, olmayan, yineleme, guven)
        satirlar.append([
            etiket[anahtar], tr(_pr(yok)),
            "-" if f is None else tri(f["Mean_Difference"]),
            "-" if f is None else f"[{tri(f['CI_Lo'])}, {tri(f['CI_Hi'])}]",
            "Katkısı yok, kaldırıldı" if anahtar == "diversity"
            else "Aşama gerekli"])
    return (["Konfigürasyon", "PR-AUC", "Çıkarılan aşamanın katkısı",
             "%95 aralık", "Karar"], satirlar)


# --- Çizelge 5.8: düğüm sayısı ---------------------------------------------

def cizelge_5_8():
    d = _ozet()
    temel = _ort(d, "baseline", "Total_Comm_MB")
    satirlar = []
    for dugum, ad in [(10, "nodes-10"), (20, "nodes-20"), (50, "baseline")]:
        toplam = _ort(d, ad, "Total_Comm_MB")
        satirlar.append([str(dugum), str(1000 // dugum),
                         tr(_pr(ad)), tr(_prsd(ad)),
                         tr(toplam, 1),
                         "referans" if dugum == 50
                         else yuzde(100 * (1 - toplam / temel))])
    return (["Düğüm sayısı", "İstemci başına kullanıcı", "PR-AUC", "sd",
             "Toplam (MB)", "Tasarruf"], satirlar)


# --- Çizelge 5.9: heterojenlik ---------------------------------------------

HETEROJENLIK = [("baseline", "IID", "FedAvg"),
                ("non-iid-baseline", "Dirichlet (αD = 0,5)", "FedAvg"),
                ("role-non-iid-baseline", "Role göre", "FedAvg"),
                ("fedprox-baseline", "IID", "FedProx"),
                ("fedprox-non-iid", "Dirichlet (αD = 0,5)", "FedProx"),
                ("role-fedprox-non-iid", "Role göre", "FedProx")]


def cizelge_5_9():
    d = _ozet()
    return (["Deney", "Bölümleme", "Birleştirme", "PR-AUC", "sd"],
            [[ad, bol, bir, tr(_pr(ad)), tr(_prsd(ad))]
             for ad, bol, bir in HETEROJENLIK])


# --- Çizelge 5.10: eşiğe bağlı başarım -------------------------------------

# Çizelge 5.10 hangi kolları gösterecek, bir sunum kararıdır: eşiğe bağlı
# eksende en çok şey anlatan dokuz kol seçilmiştir. Sayılar ölçümden gelir,
# seçim ve Türkçe etiketler buradan.
ESIK_SATIRLARI = [
    ("baseline", "baseline"),
    ("delta-0.05-downlink-fp16", "delta-0.05-downlink-fp16"),
    ("nodes-10", "nodes-10"),
    ("non-iid-baseline", "non-iid-baseline"),
    ("top-k-0.1-quant-fp16", "top-k-0.1-quant-fp16"),
    ("top-k-0.05", "top-k-0.05"),
    ("no-bottleneck", "no-bottleneck"),
    ("ablation-no-zscore", "Z-skoru yok"),
    ("ablation-no-persistence", "Zamansal kalıcılık yok"),
]


def cizelge_5_10():
    import statistics
    from config_manager import config
    from threshold_metrics import topla
    pencere = config.get("evaluation", "plateau_window_rounds")
    sonuc = topla(pencere, {kol for kol, _ in ESIK_SATIRLARI})
    satirlar = []
    for kol, etiket in ESIK_SATIRLARI:
        if kol not in sonuc or not sonuc[kol]["Max-F1"]:
            raise SystemExit(f"'{kol}' için eşiğe bağlı metrik yok.")
        o = {m: statistics.mean(v) for m, v in sonuc[kol].items() if v}
        satirlar.append([
            etiket, tr(o["Precision"]), tr(o["Recall"]), tr(o["Max-F1"]),
            tr(o["Balanced_Accuracy"]),
            f"{tr(o['TP'], 1)} / {tr(o['FP'], 1)} / {tr(o['FN'], 1)}"])
    return (["Deney", "Kesinlik", "Duyarlılık", "F1", "Dengeli doğruluk",
             "TP / FP / FN"], satirlar)


URETICILER = {
    "4.1": cizelge_4_1,
    "4.2m": cizelge_4_2m,
    "4.3": cizelge_4_3,
    "5.3": cizelge_5_3,
    "5.4": cizelge_5_4,
    "5.5": cizelge_5_5,
    "5.6": cizelge_5_6,
    "5.7": cizelge_5_7,
    "5.8": cizelge_5_8,
    "5.9": cizelge_5_9,
    "5.10": cizelge_5_10,
}


# --- Belgeyle karşılaştırma -------------------------------------------------

def belgedeki_cizelgeler(yol):
    """Belgedeki çizelgeleri numarasına göre döndürür.

    Eşleme altyazıya dayanır: bir çizelge nesnesi, kendisinden hemen önce gelen
    `Çizelge N:` altyazısına aittir. Sıraya güvenmek yanlış olurdu, çünkü kapak
    sayfasındaki düzen tabloları da nesne sayılır.
    """
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(yol)
    def metin(p):
        return "".join(n.text or "" for n in p._p.iter(qn("w:t")))

    sirali, son_altyazi = {}, None
    tablo_indeks = 0
    for ogesi in doc.element.body.iterchildren():
        etiket = ogesi.tag.split("}")[1]
        if etiket == "p":
            for p in doc.paragraphs:
                if p._p is ogesi:
                    m = re.match(r"\s*Çizelge\s+(\d+\.\d+)\s*:", metin(p))
                    if m and not p.style.name.lower().replace(" ", "").startswith(
                            ("toc", "tableoffigures")):
                        son_altyazi = m.group(1)
                    break
        elif etiket == "tbl":
            if son_altyazi is not None:
                sirali[son_altyazi] = doc.tables[tablo_indeks]
                son_altyazi = None
            tablo_indeks += 1
    return sirali


def karsilastir(no, basliklar, satirlar, tablo):
    """Üretilen ile belgedeki arasındaki farkları döndürür."""
    beklenen = [basliklar] + satirlar
    mevcut = [[h.text.strip() for h in r.cells] for r in tablo.rows]
    fark = []
    if len(beklenen) != len(mevcut):
        fark.append(f"satır sayısı: üretilen {len(beklenen)}, belgede {len(mevcut)}")
    for i, (b, m) in enumerate(zip(beklenen, mevcut)):
        if len(b) > len(m):
            fark.append(f"satır {i}: sütun sayısı {len(b)} vs {len(m)}")
            continue
        # Belgedeki fazla sütunlar elle yazılmış olabilir (örnekler, açıklama);
        # üretilen sütun sayısı kadar karşılaştırılır.
        for j, (x, y) in enumerate(zip(b, m)):
            if x.strip() != y.strip():
                fark.append(f"satır {i} sütun {j}: üretilen {x!r}, belgede {y!r}")
    return fark


def main():
    ayr = argparse.ArgumentParser(description=__doc__)
    ayr.add_argument("--out", default=VARSAYILAN_CIKTI)
    ayr.add_argument("--only", default="")
    ayr.add_argument("--check", action="store_true",
                     help="yalnız belgeyle karşılaştır, dosya yazma")
    ayr.add_argument("--iterations", type=int, default=None,
                     help="varsayılan: yapılandırmadaki bootstrap_iterations")
    a = ayr.parse_args()

    os.chdir(KOK)
    istenen = [x.strip() for x in a.only.split(",") if x.strip()] or list(URETICILER)
    bilinmeyen = [x for x in istenen if x not in URETICILER]
    if bilinmeyen:
        raise SystemExit(f"Bilinmeyen çizelge: {bilinmeyen}. "
                         f"Üretilebilenler: {sorted(URETICILER)}")

    if not a.check:
        os.makedirs(a.out, exist_ok=True)

    belgedeki = {}
    if os.path.exists(BELGE):
        try:
            belgedeki = belgedeki_cizelgeler(BELGE)
        except Exception as hata:
            print(f"!  belge okunamadı ({hata}); karşılaştırma atlanıyor")

    inceleme = ["# Üretilen çizelgeler\n"]
    sorunlu = []
    for no in istenen:
        try:
            sonuc = (URETICILER[no](yineleme=a.iterations)
                     if no in ("5.6", "5.7") else URETICILER[no]())
        except SystemExit as hata:
            print(f"!  Çizelge {no}: {hata}")
            sorunlu.append(no)
            continue
        basliklar, satirlar = sonuc

        if not a.check:
            with open(os.path.join(a.out, f"cizelge_{no}.tsv"), "w",
                      encoding="utf-8") as f:
                f.write("\t".join(basliklar) + "\n")
                for s in satirlar:
                    f.write("\t".join(s) + "\n")

        inceleme.append(f"\n## Çizelge {no}\n")
        inceleme.append("| " + " | ".join(basliklar) + " |")
        inceleme.append("|" + "---|" * len(basliklar))
        for s in satirlar:
            inceleme.append("| " + " | ".join(s) + " |")

        if no in belgedeki:
            fark = karsilastir(no, basliklar, satirlar, belgedeki[no])
            if fark:
                print(f"!  Çizelge {no}: belgeden {len(fark)} farkı var")
                for f in fark[:8]:
                    print(f"     {f}")
                if len(fark) > 8:
                    print(f"     ... {len(fark) - 8} fark daha")
                inceleme.append(f"\n**Belgeden farkı: {len(fark)} hücre.**\n")
            else:
                print(f"   Çizelge {no}: belgeyle birebir")
        elif no != "4.2m":
            print(f"   Çizelge {no}: üretildi (belgede eşleşen çizelge bulunamadı)")

    if not a.check:
        yol = os.path.join(a.out, "hepsi.md")
        with open(yol, "w", encoding="utf-8") as f:
            f.write("\n".join(inceleme) + "\n")
        print(f"\nYazıldı: {a.out}/  ({len(istenen) - len(sorunlu)} çizelge, "
              f"sekme ayraçlı + hepsi.md)")

    print("\nElle bakımda kalanlar:")
    for no, sebep in sorted(ELLE.items()):
        print(f"  Çizelge {no}: {sebep}")
    return 1 if sorunlu else 0


if __name__ == "__main__":
    sys.exit(main())
